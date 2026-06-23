from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXCEL_DIR = ROOT / "data" / "excel"
CSV_DIR = ROOT / "data" / "csv"

WORKBOOKS = [
    {
        "path": EXCEL_DIR / "review_defect_statistics.xlsx",
        "csv_folder": "review_defect_statistics",
    },
    {
        "path": EXCEL_DIR / "study_coding_workbook.xlsx",
        "csv_folder": "study_coding_workbook",
    },
]


def clean_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]+', "_", name)
    name = re.sub(r"\s+", "_", name.strip())
    return name


def cell_to_text(value) -> str:
    return "" if value is None else str(value)


def export_workbook(workbook_path: Path, csv_folder: str) -> int:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Missing workbook: {workbook_path}")

    output_dir = CSV_DIR / csv_folder
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_count = 0
        for worksheet in workbook.worksheets:
            output_path = output_dir / f"{clean_filename(worksheet.title)}.csv"
            with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, lineterminator="\n")
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=worksheet.max_row,
                    min_col=1,
                    max_col=worksheet.max_column,
                    values_only=True,
                ):
                    writer.writerow([cell_to_text(value) for value in row])
            sheet_count += 1
        return sheet_count
    finally:
        workbook.close()


def main() -> None:
    total = 0
    for workbook in WORKBOOKS:
        count = export_workbook(workbook["path"], workbook["csv_folder"])
        total += count
        print(f"Exported {count} sheets from {workbook['path'].name}")
    print(f"Done. Exported {total} CSV files to {CSV_DIR}")


if __name__ == "__main__":
    main()
